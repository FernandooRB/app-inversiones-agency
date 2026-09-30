"""Private, read-only preflight for one observed GBM portfolio XLSX layout.

The result contains counts and a fingerprint, never positions, amounts or labels.
It does not prove the workbook's account, valuation date or movement coverage.
"""

from __future__ import annotations

import argparse
import json
import re
import unicodedata
import warnings
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

MAX_XLSX_BYTES = 2_000_000
MAX_UNCOMPRESSED_BYTES = 15_000_000
MAX_ZIP_ENTRIES = 200
MAX_ROWS = 500
MAX_COLUMNS = 11
NUMBER = re.compile(r"-?\$?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?\Z")


class PortfolioXlsxError(ValueError):
    """A format error whose message does not repeat private workbook content."""


def _fold(value: object) -> str:
    text = str(value or "").strip().lower()
    return "".join(char for char in unicodedata.normalize("NFKD", text)
                   if not unicodedata.combining(char))


def _number(value: object, *, positive: bool = False) -> Decimal:
    if value is None or isinstance(value, bool):
        raise PortfolioXlsxError("Falta un dato numérico requerido en la cartera GBM.")
    token = str(value).strip().replace(" ", "")
    if not NUMBER.fullmatch(token):
        raise PortfolioXlsxError("Un dato numérico tiene formato no reconocido.")
    try:
        result = Decimal(token.replace("$", "").replace(",", ""))
    except InvalidOperation as exc:
        raise PortfolioXlsxError("Un dato numérico no es válido.") from exc
    if result < 0 or (positive and result == 0):
        raise PortfolioXlsxError("Una cantidad o valuación está fuera del alcance admitido.")
    return result


def _validate_zip(raw: bytes) -> None:
    try:
        with ZipFile(BytesIO(raw)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_ENTRIES or sum(
                item.file_size for item in entries
            ) > MAX_UNCOMPRESSED_BYTES:
                raise PortfolioXlsxError("El libro supera el límite de contenido admitido.")
            paths = {item.filename.lower() for item in entries}
            if len(paths) != len(entries) or "xl/workbook.xml" not in paths or any(
                "vbaproject" in path or "externallinks/" in path for path in paths
            ):
                raise PortfolioXlsxError("El libro contiene partes no admitidas.")
    except BadZipFile as exc:
        raise PortfolioXlsxError("El archivo no es un XLSX válido.") from exc


def _row_has_values(row: tuple[object, ...]) -> bool:
    return any(cell.value is not None and str(cell.value).strip() for cell in row)


def _validate_header(row: tuple[object, ...]) -> None:
    expected = {
        0: "emisora/fondo", 2: "costo promedio", 3: "precio mercado",
        5: "valor mercado", 6: "p / m", 7: "% var. hist.",
        8: "% var. dia.", 9: "imp x cto", 10: "% cartera",
    }
    if any(_fold(row[index].value) != value for index, value in expected.items()):
        raise PortfolioXlsxError("Los encabezados no corresponden al diseño GBM observado.")
    if not str(row[1].value or "").strip():
        raise PortfolioXlsxError("Falta el encabezado de cantidad de títulos.")


def _validate_position(row: tuple[object, ...], seen: set[str], *, cash: bool) -> None:
    label = _fold(row[0].value)
    if not label or label in seen:
        raise PortfolioXlsxError("Hay una posición sin etiqueta o duplicada.")
    seen.add(label)
    _number(row[1].value, positive=not cash)
    _number(row[3].value)
    _number(row[5].value)


def inspect_portfolio_xlsx(raw: bytes) -> dict[str, object]:
    """Validate structure only; return no user-controlled text from the workbook."""
    if not 0 < len(raw) <= MAX_XLSX_BYTES:
        raise PortfolioXlsxError("El XLSX está vacío o excede 2 MB.")
    try:
        _validate_zip(raw)
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", message="Cannot parse header or footer")
            book = load_workbook(BytesIO(raw), read_only=True, data_only=False,
                                 keep_links=False)
        try:
            if len(book.worksheets) != 1:
                raise PortfolioXlsxError("El libro debe tener una sola hoja.")
            sheet = book.worksheets[0]
            if not 4 <= sheet.max_row <= MAX_ROWS or sheet.max_column != MAX_COLUMNS:
                raise PortfolioXlsxError("El tamaño de la hoja no corresponde al diseño admitido.")
            with warnings.catch_warnings():
                warnings.filterwarnings("ignore", message="Cannot parse header or footer")
                rows = list(sheet.iter_rows())
            if any(cell.data_type == "f" for row in rows for cell in row):
                raise PortfolioXlsxError("El libro contiene fórmulas; requiere revisión manual.")
            if _fold(rows[0][0].value) != "mercado de capitales nacional" or any(
                cell.value is not None for cell in rows[0][1:]
            ):
                raise PortfolioXlsxError("La primera sección no corresponde al diseño observado.")
            _validate_header(rows[1])
            equity_labels: set[str] = set()
            cash_labels: set[str] = set()
            index = 2
            while index < len(rows) and _fold(rows[index][0].value) != "efectivo":
                if _row_has_values(rows[index]):
                    _validate_position(rows[index], equity_labels, cash=False)
                index += 1
            if not equity_labels or index + 1 >= len(rows):
                raise PortfolioXlsxError("Falta la sección de posiciones o efectivo.")
            if any(cell.value is not None for cell in rows[index][1:]):
                raise PortfolioXlsxError("La sección de efectivo contiene columnas inesperadas.")
            _validate_header(rows[index + 1])
            for row in rows[index + 2:]:
                if _row_has_values(row):
                    _validate_position(row, cash_labels, cash=True)
            if not cash_labels:
                raise PortfolioXlsxError("La sección de efectivo está vacía.")
            return {
                "status": "ESTRUCTURA_PLAUSIBLE_NO_AUTENTICADA",
                "sha256": sha256(raw).hexdigest(),
                "position_rows": len(equity_labels),
                "cash_rows": len(cash_labels),
                "contains_movements": False,
                "as_of_date_verified": False,
                "account_verified": False,
                "caution": "Sólo valida el formato observado. No acredita fecha, cuenta, "
                           "completitud, valuación, precios ni movimientos.",
            }
        finally:
            book.close()
    except PortfolioXlsxError:
        raise
    except Exception as exc:
        raise PortfolioXlsxError("No se pudo leer el XLSX de cartera GBM.") from exc


def main() -> int:
    parser = argparse.ArgumentParser(description="Preflight privado de cartera GBM XLSX")
    parser.add_argument("file", type=Path)
    args = parser.parse_args()
    try:
        if not args.file.is_file() or args.file.is_symlink():
            raise PortfolioXlsxError("Selecciona un archivo XLSX local válido.")
        result = inspect_portfolio_xlsx(args.file.read_bytes())
    except (OSError, PortfolioXlsxError) as exc:
        result = {"status": "REVIEW_REQUIRED", "error": (
            str(exc) if isinstance(exc, PortfolioXlsxError)
            else "No se pudo leer el archivo local."
        )}
    print(json.dumps(result, ensure_ascii=False, sort_keys=True))
    return 0 if result["status"] == "ESTRUCTURA_PLAUSIBLE_NO_AUTENTICADA" else 2


if __name__ == "__main__":
    raise SystemExit(main())
